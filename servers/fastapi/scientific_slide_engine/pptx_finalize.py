from __future__ import annotations

import hashlib
import os
import tempfile
import zipfile
from pathlib import Path
from typing import Sequence

from pptx import Presentation


class ScientificPptxFinalizationError(RuntimeError):
    pass


def _slide_xml_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        return {
            name: hashlib.sha256(zf.read(name)).hexdigest()
            for name in zf.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        }


def _norm_note(value: str | None) -> str:
    return (value or "").replace("\r\n", "\n").replace("\r", "\n")


def finalize_scientific_pptx(
    pptx_path: str | os.PathLike[str],
    speaker_notes: Sequence[str],
) -> dict[str, object]:
    """Inject exact notes into a PPTX already rendered by Presenton/export-core.

    The post-processor is deliberately forbidden from changing slide canvas XML.
    Presenton remains the renderer; this step only completes/validates notesSlides.
    """
    path = Path(pptx_path)
    if not path.is_file():
        raise ScientificPptxFinalizationError(f"PPTX not found: {path}")

    before_hashes = _slide_xml_hashes(path)
    prs = Presentation(str(path))
    if len(prs.slides) != len(speaker_notes):
        raise ScientificPptxFinalizationError(
            f"PPTX slide count {len(prs.slides)} != notes count {len(speaker_notes)}"
        )

    for slide, note in zip(prs.slides, speaker_notes):
        slide.notes_slide.notes_text_frame.text = _norm_note(note)

    with tempfile.NamedTemporaryFile(
        dir=str(path.parent),
        prefix=f".{path.stem}.scientific-notes-",
        suffix=".pptx",
        delete=False,
    ) as handle:
        temp_path = Path(handle.name)

    try:
        prs.save(str(temp_path))
        after_hashes = _slide_xml_hashes(temp_path)
        if before_hashes != after_hashes:
            raise ScientificPptxFinalizationError(
                "speaker-note finalization modified slide canvas XML; refusing export"
            )

        check = Presentation(str(temp_path))
        actual_notes = [
            _norm_note(slide.notes_slide.notes_text_frame.text)
            for slide in check.slides
        ]
        expected_notes = [_norm_note(note) for note in speaker_notes]
        if actual_notes != expected_notes:
            mismatch = next(
                (
                    index
                    for index, (actual, expected) in enumerate(
                        zip(actual_notes, expected_notes), 1
                    )
                    if actual != expected
                ),
                None,
            )
            raise ScientificPptxFinalizationError(
                f"speaker-note round-trip mismatch at slide {mismatch}"
            )

        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    with zipfile.ZipFile(path) as zf:
        note_parts = [
            name
            for name in zf.namelist()
            if name.startswith("ppt/notesSlides/notesSlide")
            and name.endswith(".xml")
        ]
        slide_parts = [
            name
            for name in zf.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        editable_text_runs = sum(
            zf.read(name).count(b"<a:t>") for name in slide_parts
        )

    return {
        "status": "PASS",
        "slide_count": len(before_hashes),
        "notes_count": len(note_parts),
        "notes_exact": True,
        "slide_canvas_unchanged": True,
        "editable_text_runs": editable_text_runs,
    }


__all__ = ["ScientificPptxFinalizationError", "finalize_scientific_pptx"]
