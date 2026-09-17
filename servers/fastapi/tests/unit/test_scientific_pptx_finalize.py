from pathlib import Path

import pytest
from pptx import Presentation

from scientific_slide_engine.pptx_finalize import (
    ScientificPptxFinalizationError,
    finalize_scientific_pptx,
)


def _make_editable_deck(path: Path, slide_count: int = 2) -> None:
    prs = Presentation()
    for index in range(slide_count):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(1000000, 1000000, 3000000, 1000000)
        box.text = f"Editable {index + 1}"
    prs.save(path)


def test_exact_notes_are_injected_without_modifying_slide_canvas(tmp_path):
    path = tmp_path / "deck.pptx"
    _make_editable_deck(path)
    notes = ["Nota uno", "Nota due\nSeconda riga"]

    report = finalize_scientific_pptx(path, notes)

    assert report["status"] == "PASS"
    assert report["notes_count"] == 2
    assert report["slide_canvas_unchanged"] is True
    assert report["editable_text_runs"] >= 2

    prs = Presentation(path)
    assert [slide.notes_slide.notes_text_frame.text for slide in prs.slides] == notes
    assert any(
        shape.text == "Editable 1"
        for shape in prs.slides[0].shapes
        if hasattr(shape, "text")
    )


def test_notes_count_mismatch_is_rejected(tmp_path):
    path = tmp_path / "deck.pptx"
    _make_editable_deck(path, 2)

    with pytest.raises(ScientificPptxFinalizationError):
        finalize_scientific_pptx(path, ["only one note"])
